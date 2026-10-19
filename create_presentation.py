import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def build_presentation(output_path):
    prs = Presentation()
    # 16:9 widescreen dimensions
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    blank_layout = prs.slide_layouts[6] # completely blank layout

    # Color Palette
    PRIMARY = RGBColor(31, 93, 76)       # #1F5D4C - Deep Campus Green
    DARK = RGBColor(20, 50, 42)          # #14322A - Dark Forest
    LIGHT_BG = RGBColor(246, 249, 246)   # #F6F9F6 - Soft Mint Off-White
    WHITE = RGBColor(255, 255, 255)
    CARD_BG = RGBColor(255, 255, 255)
    BORDER = RGBColor(215, 226, 218)
    TEXT_DARK = RGBColor(25, 38, 33)
    TEXT_MUTED = RGBColor(100, 116, 108)
    ACCENT = RGBColor(200, 70, 75)       # Reddish Coral
    GOLD = RGBColor(218, 145, 0)
    TINT_GREEN = RGBColor(235, 243, 238)

    def set_slide_background(slide, color=LIGHT_BG):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = color
        bg.line.fill.background()
        return bg

    def add_header(slide, title_text, category="CAMPUS MARKETPLACE — UNIVERSITY TEAM PROJECT"):
        # Category Badge
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11), Inches(0.4))
        tf = cat_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = category.upper()
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = PRIMARY
        
        # Title
        t_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.5), Inches(0.7))
        tf2 = t_box.text_frame
        tf2.word_wrap = True
        p2 = tf2.paragraphs[0]
        p2.text = title_text
        p2.font.size = Pt(22)
        p2.font.bold = True
        p2.font.color.rgb = DARK

        # Header underline
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.45), Inches(11.733), Inches(0.04))
        line.fill.solid()
        line.fill.fore_color.rgb = PRIMARY
        line.line.fill.background()

    # ==========================================================
    # SLIDE 1: Title Slide with Member Details Placeholders
    # ==========================================================
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1, DARK)

    # Hero Banner Card
    banner = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.8), Inches(11.733), Inches(2.5))
    banner.fill.solid()
    banner.fill.fore_color.rgb = PRIMARY
    banner.line.color.rgb = RGBColor(60, 140, 115)
    banner.line.width = Pt(1.5)

    btf = banner.text_frame
    btf.word_wrap = True
    bp1 = btf.paragraphs[0]
    bp1.text = "🎓 UNIVERSITY SOFTWARE ENGINEERING PROJECT"
    bp1.font.size = Pt(11)
    bp1.font.bold = True
    bp1.font.color.rgb = RGBColor(165, 220, 195)

    bp2 = btf.add_paragraph()
    bp2.text = "Campus Marketplace"
    bp2.font.size = Pt(36)
    bp2.font.bold = True
    bp2.font.color.rgb = WHITE

    bp3 = btf.add_paragraph()
    bp3.text = "A Hyper-Localized Peer-to-Peer Trading Platform for Sri Lankan University Students"
    bp3.font.size = Pt(14)
    bp3.font.color.rgb = RGBColor(225, 240, 232)

    bp4 = btf.add_paragraph()
    bp4.text = "Tech Stack: React Native (Expo SDK 57) · Google Cloud Firestore & Storage · TypeScript"
    bp4.font.size = Pt(11)
    bp4.font.color.rgb = RGBColor(170, 210, 190)

    # Member Placeholders Grid Title
    lbl = s1.shapes.add_textbox(Inches(0.8), Inches(3.5), Inches(11.7), Inches(0.4))
    lbl_tf = lbl.text_frame
    lbl_p = lbl_tf.paragraphs[0]
    lbl_p.text = "👥 PROJECT TEAM MEMBERS & ROLES (PLACEHOLDER SECTION)"
    lbl_p.font.size = Pt(11)
    lbl_p.font.bold = True
    lbl_p.font.color.rgb = RGBColor(170, 220, 195)

    # 5 Member Detail Cards
    card_w = Inches(2.22)
    card_h = Inches(3.0)
    gap = Inches(0.15)
    left_start = Inches(0.8)

    members_data = [
        ("Member 1", "Campus Search & Request Board (WTB)", "[STUDENT 1 NAME]", "[STUDENT 1 ID]"),
        ("Member 2", "Live Chat & Price Bargaining", "[STUDENT 2 NAME]", "[STUDENT 2 ID]"),
        ("Member 3", "Listing CRUD & Safe Meetup", "[STUDENT 3 NAME]", "[STUDENT 3 ID]"),
        ("Member 4", "Auth Security & Seller Reviews", "[STUDENT 4 NAME]", "[STUDENT 4 ID]"),
        ("Member 5", "Data Models & Scam Moderation", "[STUDENT 5 NAME]", "[STUDENT 5 ID]"),
    ]

    for i, (m_tag, role, name_ph, id_ph) in enumerate(members_data):
        c_left = left_start + i * (card_w + gap)
        c = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c_left, Inches(3.9), card_w, card_h)
        c.fill.solid()
        c.fill.fore_color.rgb = RGBColor(28, 65, 55)
        c.line.color.rgb = RGBColor(60, 120, 100)
        c.line.width = Pt(1)

        ctf = c.text_frame
        ctf.word_wrap = True
        
        cp1 = ctf.paragraphs[0]
        cp1.text = m_tag.upper()
        cp1.font.size = Pt(10)
        cp1.font.bold = True
        cp1.font.color.rgb = GOLD

        cp2 = ctf.add_paragraph()
        cp2.text = name_ph
        cp2.font.size = Pt(13)
        cp2.font.bold = True
        cp2.font.color.rgb = WHITE

        cp3 = ctf.add_paragraph()
        cp3.text = f"ID: {id_ph}"
        cp3.font.size = Pt(11)
        cp3.font.color.rgb = RGBColor(190, 215, 205)

        cp4 = ctf.add_paragraph()
        cp4.text = f"\nAssigned Feature:\n{role}"
        cp4.font.size = Pt(10)
        cp4.font.color.rgb = RGBColor(220, 240, 230)

    # ==========================================================
    # SLIDE 2: Team Members & Contribution Matrix (Academic View)
    # ==========================================================
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_background(s2)
    add_header(s2, "Team Workload Distribution & Feature Matrix (20% Equal Each)")

    matrix_data = [
        ("MEMBER", "NAME & STUDENT ID", "ASSIGNED SUBSYSTEM", "SOURCE CODE FILES", "WEIGHT"),
        ("Member 1", "[Member 1 Name - ID: IT21XXXXXX]", "Campus Search, Price Filter & Request Board (WTB)", "ItemRequestsModal.tsx, ExplorePage.tsx, data.ts", "20%"),
        ("Member 2", "[Member 2 Name - ID: IT21XXXXXX]", "Real-Time Firestore Chat & Price Bargaining Engine", "MakeOfferModal.tsx, MessagesPage.tsx", "20%"),
        ("Member 3", "[Member 3 Name - ID: IT21XXXXXX]", "Seller Inventory CRUD & Safe Campus Meetup Coordinator", "MeetupSchedulerModal.tsx, MyListingsPage.tsx, ListingCard.tsx", "20%"),
        ("Member 4", "[Member 4 Name - ID: IT21XXXXXX]", "User Authentication Security & Seller Reputation System", "SellerReviewsModal.tsx, ProfilePage.tsx", "20%"),
        ("Member 5", "[Member 5 Name - ID: IT21XXXXXX]", "TypeScript Data Architecture & Scam Moderation Shield", "ReportListingModal.tsx, SavedPage.tsx, types.ts", "20%"),
    ]

    # Render Table
    rows = len(matrix_data)
    cols = 5
    t_left = Inches(0.8)
    t_top = Inches(1.8)
    t_width = Inches(11.733)
    t_height = Inches(4.5)
    
    table_shape = s2.shapes.add_table(rows, cols, t_left, t_top, t_width, t_height)
    tbl = table_shape.table
    tbl.columns[0].width = Inches(1.4)
    tbl.columns[1].width = Inches(2.8)
    tbl.columns[2].width = Inches(3.6)
    tbl.columns[3].width = Inches(3.1)
    tbl.columns[4].width = Inches(0.833)

    for r_idx, row in enumerate(matrix_data):
        for c_idx, val in enumerate(row):
            cell = tbl.cell(r_idx, c_idx)
            cell.text = val
            cell_p = cell.text_frame.paragraphs[0]
            if r_idx == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = PRIMARY
                cell_p.font.bold = True
                cell_p.font.size = Pt(11)
                cell_p.font.color.rgb = WHITE
            else:
                cell.fill.solid()
                cell.fill.fore_color.rgb = CARD_BG if r_idx % 2 == 1 else TINT_GREEN
                cell_p.font.size = Pt(10)
                cell_p.font.color.rgb = TEXT_DARK
                if c_idx == 4:
                    cell_p.font.bold = True
                    cell_p.font.color.rgb = PRIMARY

    # ==========================================================
    # SLIDE 3: Problem Statement & Motivation
    # ==========================================================
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_background(s3)
    add_header(s3, "Problem Statement & Motivation")

    problems = [
        ("💸 High Academic Costs", "University students spend thousands of LKR each semester on mandatory textbooks, senior lecture notes, and drafting tools that are only used for a single semester.", ACCENT),
        ("📍 Lack of Campus Localization", "General platforms (Facebook Groups, Ikman) mix listings island-wide. A student in Colombo cannot easily filter items held by someone in the Moratuwa or Peradeniya campus.", PRIMARY),
        ("🛡️ Safety Risks & Scams", "Meeting strangers off-campus carries safety hazards. Furthermore, buyers frequently face advance payment scams, counterfeit past papers, and unverified sellers.", GOLD),
    ]

    for i, (p_title, p_desc, bar_col) in enumerate(problems):
        card_top = Inches(1.9 + i * 1.6)
        c = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), card_top, Inches(11.733), Inches(1.35))
        c.fill.solid()
        c.fill.fore_color.rgb = CARD_BG
        c.line.color.rgb = BORDER
        c.line.width = Pt(1)

        # Left color bar
        bar = s3.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), card_top, Inches(0.18), Inches(1.35))
        bar.fill.solid()
        bar.fill.fore_color.rgb = bar_col
        bar.line.fill.background()

        ctf = c.text_frame
        ctf.word_wrap = True
        cp1 = ctf.paragraphs[0]
        cp1.text = f"   {p_title}"
        cp1.font.size = Pt(15)
        cp1.font.bold = True
        cp1.font.color.rgb = DARK

        cp2 = ctf.add_paragraph()
        cp2.text = f"   {p_desc}"
        cp2.font.size = Pt(11.5)
        cp2.font.color.rgb = TEXT_MUTED

    # ==========================================================
    # SLIDE 4: Proposed Solution & Core System Capabilities
    # ==========================================================
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_background(s4)
    add_header(s4, "Proposed Solution: Hyper-Localized Campus Marketplace")

    pillars = [
        ("1. Campus Discovery", "Filtered by 17 state universities & LKR price ranges. Features a Student Request Board (WTB) for unlisted goods.", Inches(0.8), Inches(1.8)),
        ("2. Real-Time Chat & Offers", "Structured 'Make an Offer' engine with automated percentage discounts & real-time Firestore messaging.", Inches(6.8), Inches(1.8)),
        ("3. Safe Meetup Logistics", "Designated physical safe exchange zones (Library, Canteen, Security Gate) with scheduled time windows.", Inches(0.8), Inches(4.3)),
        ("4. Trust, Auth & Safety", "Email/Password security, verified 5-star seller peer reviews, trust badges, and scam report shield.", Inches(6.8), Inches(4.3)),
    ]

    for title, desc, l, t in pillars:
        box = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, Inches(5.7), Inches(2.1))
        box.fill.solid()
        box.fill.fore_color.rgb = CARD_BG
        box.line.color.rgb = BORDER
        box.line.width = Pt(1)

        # Inner header accent
        hb = s4.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, Inches(5.7), Inches(0.55))
        hb.fill.solid()
        hb.fill.fore_color.rgb = TINT_GREEN
        hb.line.fill.background()

        hbtf = hb.text_frame
        hbp = hbtf.paragraphs[0]
        hbp.text = f"  {title}"
        hbp.font.size = Pt(13)
        hbp.font.bold = True
        hbp.font.color.rgb = PRIMARY

        btf = box.text_frame
        btf.word_wrap = True
        bp1 = btf.paragraphs[0]
        bp1.text = "\n\n" + desc
        bp1.font.size = Pt(12)
        bp1.font.color.rgb = TEXT_DARK

    # ==========================================================
    # SLIDE 5: Technical Architecture & System Stack
    # ==========================================================
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_background(s5)
    add_header(s5, "System Architecture & Engineering Stack")

    tech_cols = [
        ("FRONTEND", "React Native & Expo SDK 57", [
            "• Universal iOS, Android, and Web compatibility",
            "• TypeScript 6.0 with strict data contracts",
            "• Custom design system (#1F5D4C Campus Green)",
            "• Single-Modal view routing engine"
        ], Inches(0.8)),
        ("BACKEND", "Firebase Cloud Services", [
            "• Firebase Authentication (Email/Password & Google)",
            "• Cloud Firestore NoSQL Database",
            "• Real-time snapshot listeners for live chat",
            "• Firebase Cloud Storage for photo uploads"
        ], Inches(4.8)),
        ("SECURITY & DEPLOY", "Rules & Quality Assurance", [
            "• Firestore Security Rules: seller-owned edits",
            "• Private message subcollections access control",
            "• Zero TypeScript compiler errors (npx tsc)",
            "• Modular isolation under src/features/"
        ], Inches(8.8)),
    ]

    for title, sub, points, l in tech_cols:
        col_box = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, Inches(1.8), Inches(3.7), Inches(4.8))
        col_box.fill.solid()
        col_box.fill.fore_color.rgb = CARD_BG
        col_box.line.color.rgb = BORDER
        col_box.line.width = Pt(1)

        c_tf = col_box.text_frame
        c_tf.word_wrap = True
        cp1 = c_tf.paragraphs[0]
        cp1.text = title
        cp1.font.size = Pt(11)
        cp1.font.bold = True
        cp1.font.color.rgb = PRIMARY

        cp2 = c_tf.add_paragraph()
        cp2.text = sub
        cp2.font.size = Pt(13)
        cp2.font.bold = True
        cp2.font.color.rgb = DARK

        for p in points:
            cp = c_tf.add_paragraph()
            cp.text = p
            cp.font.size = Pt(11)
            cp.font.color.rgb = TEXT_MUTED

    # ==========================================================
    # SLIDE 6: Member 1 Module Deep Dive
    # ==========================================================
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_background(s6)
    add_header(s6, "Module 1: Campus Search & Student Request Board (WTB)")

    # Left: Developer Details Card
    dev1 = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(4.2), Inches(4.8))
    dev1.fill.solid()
    dev1.fill.fore_color.rgb = TINT_GREEN
    dev1.line.color.rgb = PRIMARY
    dev1.line.width = Pt(1)

    d1_tf = dev1.text_frame
    d1_tf.word_wrap = True
    d1_p = d1_tf.paragraphs[0]
    d1_p.text = "👤 LEAD DEVELOPER: MEMBER 1"
    d1_p.font.size = Pt(12)
    d1_p.font.bold = True
    d1_p.font.color.rgb = PRIMARY

    dp2 = d1_tf.add_paragraph()
    dp2.text = "Student Name: [MEMBER 1 NAME]\nStudent ID: [MEMBER 1 ID]\nContribution Weight: 20%"
    dp2.font.size = Pt(12)
    dp2.font.color.rgb = DARK

    dp3 = d1_tf.add_paragraph()
    dp3.text = "\nDedicated Files:\n• src/features/ItemRequestsModal.tsx\n• src/pages/ExplorePage.tsx\n• src/data.ts\n\nGit Branch: feature/item-requests"
    dp3.font.size = Pt(10.5)
    dp3.font.color.rgb = TEXT_MUTED

    # Right: Technical Delivery
    r1 = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(5.3), Inches(1.8), Inches(7.2), Inches(4.8))
    r1.fill.solid()
    r1.fill.fore_color.rgb = CARD_BG
    r1.line.color.rgb = BORDER
    r1_tf = r1.text_frame
    r1_tf.word_wrap = True

    rp1 = r1_tf.paragraphs[0]
    rp1.text = "Key Functional & Technical Deliverables"
    rp1.font.size = Pt(15)
    rp1.font.bold = True
    rp1.font.color.rgb = DARK

    rp_body = [
        ("1. Sri Lankan State University Mapping:", "Replaced generic 'North Campus' strings with 17 state universities (Moratuwa, Colombo, Peradeniya, Sri Jayewardenepura, Kelaniya, etc.)"),
        ("2. Real-Time Price Range Filtering:", "Implemented Min/Max LKR sliders to ensure students can search within exact budgets."),
        ("3. Student Request Board (WTB):", "Engineered ItemRequestsModal.tsx where students can post wanted requests with target budgets and urgency levels ('Urgent (Exam Soon)' vs 'Normal')."),
        ("4. Seamless Peer Contact:", "Included quick 'I have this!' response triggers connecting directly into messaging.")
    ]
    for h, b in rp_body:
        p_item = r1_tf.add_paragraph()
        p_item.text = f"\n• {h} {b}"
        p_item.font.size = Pt(11)
        p_item.font.color.rgb = TEXT_DARK

    # ==========================================================
    # SLIDE 7: Member 2 Module Deep Dive
    # ==========================================================
    s7 = prs.slides.add_slide(blank_layout)
    set_slide_background(s7)
    add_header(s7, "Module 2: Live Chat Subsystem & Price Bargaining Engine")

    dev2 = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(4.2), Inches(4.8))
    dev2.fill.solid()
    dev2.fill.fore_color.rgb = TINT_GREEN
    dev2.line.color.rgb = PRIMARY
    d2_tf = dev2.text_frame
    d2_tf.word_wrap = True
    d2_p = d2_tf.paragraphs[0]
    d2_p.text = "👤 LEAD DEVELOPER: MEMBER 2"
    d2_p.font.size = Pt(12)
    d2_p.font.bold = True
    d2_p.font.color.rgb = PRIMARY

    dp2 = d2_tf.add_paragraph()
    dp2.text = "Student Name: [MEMBER 2 NAME]\nStudent ID: [MEMBER 2 ID]\nContribution Weight: 20%"
    dp2.font.size = Pt(12)
    dp2.font.color.rgb = DARK

    dp3 = d2_tf.add_paragraph()
    dp3.text = "\nDedicated Files:\n• src/features/MakeOfferModal.tsx\n• src/pages/MessagesPage.tsx\n\nGit Branch: feature/make-offer"
    dp3.font.size = Pt(10.5)
    dp3.font.color.rgb = TEXT_MUTED

    r2 = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(5.3), Inches(1.8), Inches(7.2), Inches(4.8))
    r2.fill.solid()
    r2.fill.fore_color.rgb = CARD_BG
    r2.line.color.rgb = BORDER
    r2_tf = r2.text_frame
    r2_tf.word_wrap = True

    rp2 = r2_tf.paragraphs[0]
    rp2.text = "Key Functional & Technical Deliverables"
    rp2.font.size = Pt(15)
    rp2.font.bold = True
    rp2.font.color.rgb = DARK

    rp_body2 = [
        ("1. Real-Time Chat Engine:", "Transformed the empty 20-line placeholder into a full Cloud Firestore query pipeline with timestamped subcollections."),
        ("2. Price Bargaining (Make an Offer):", "Built MakeOfferModal.tsx with quick percentage discount chips (-5%, -10%, -15%, -20%) and custom offer inputs."),
        ("3. Instant Savings Calculation:", "Dynamically calculates and displays student savings in LKR before dispatching offers to sellers."),
        ("4. Deal Sweeteners:", "Integrated preset negotiation chips ('Can pick up today', 'Cash on delivery') allowing seamless trade terms.")
    ]
    for h, b in rp_body2:
        p_item = r2_tf.add_paragraph()
        p_item.text = f"\n• {h} {b}"
        p_item.font.size = Pt(11)
        p_item.font.color.rgb = TEXT_DARK

    # ==========================================================
    # SLIDE 8: Member 3 Module Deep Dive
    # ==========================================================
    s8 = prs.slides.add_slide(blank_layout)
    set_slide_background(s8)
    add_header(s8, "Module 3: Seller Inventory CRUD & Safe Meetup Coordinator")

    dev3 = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(4.2), Inches(4.8))
    dev3.fill.solid()
    dev3.fill.fore_color.rgb = TINT_GREEN
    dev3.line.color.rgb = PRIMARY
    d3_tf = dev3.text_frame
    d3_tf.word_wrap = True
    d3_p = d3_tf.paragraphs[0]
    d3_p.text = "👤 LEAD DEVELOPER: MEMBER 3"
    d3_p.font.size = Pt(12)
    d3_p.font.bold = True
    d3_p.font.color.rgb = PRIMARY

    dp3 = d3_tf.add_paragraph()
    dp3.text = "Student Name: [MEMBER 3 NAME]\nStudent ID: [MEMBER 3 ID]\nContribution Weight: 20%"
    dp3.font.size = Pt(12)
    dp3.font.color.rgb = DARK

    dp3 = d3_tf.add_paragraph()
    dp3.text = "\nDedicated Files:\n• src/features/MeetupSchedulerModal.tsx\n• src/pages/MyListingsPage.tsx\n• src/components/ListingCard.tsx\n\nGit Branch: feature/safe-meetup"
    dp3.font.size = Pt(10.5)
    dp3.font.color.rgb = TEXT_MUTED

    r3 = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(5.3), Inches(1.8), Inches(7.2), Inches(4.8))
    r3.fill.solid()
    r3.fill.fore_color.rgb = CARD_BG
    r3.line.color.rgb = BORDER
    r3_tf = r3.text_frame
    r3_tf.word_wrap = True

    rp3 = r3_tf.paragraphs[0]
    rp3.text = "Key Functional & Technical Deliverables"
    rp3.font.size = Pt(15)
    rp3.font.bold = True
    rp3.font.color.rgb = DARK

    rp_body3 = [
        ("1. Full Listing CRUD Flow:", "Implemented Edit listing forms with pre-filled state, and Delete listings with native confirmation alerts."),
        ("2. Cloud Photo Uploads:", "Integrated expo-image-picker and Firebase Cloud Storage blob uploads under listing-images/{uid}/{timestamp}."),
        ("3. Safe Campus Meetup Coordinator:", "Developed MeetupSchedulerModal.tsx designating verified public zones (Library Lobby, Canteen, Security Gate)."),
        ("4. Scheduling & Safety Guidelines:", "Provided date chips, time windows ('Lunch Break', 'Afternoon'), and student safety checklists.")
    ]
    for h, b in rp_body3:
        p_item = r3_tf.add_paragraph()
        p_item.text = f"\n• {h} {b}"
        p_item.font.size = Pt(11)
        p_item.font.color.rgb = TEXT_DARK

    # ==========================================================
    # SLIDE 9: Member 4 Module Deep Dive
    # ==========================================================
    s9 = prs.slides.add_slide(blank_layout)
    set_slide_background(s9)
    add_header(s9, "Module 4: Authentication Security & Seller Reputation")

    dev4 = s9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(4.2), Inches(4.8))
    dev4.fill.solid()
    dev4.fill.fore_color.rgb = TINT_GREEN
    dev4.line.color.rgb = PRIMARY
    d4_tf = dev4.text_frame
    d4_tf.word_wrap = True
    d4_p = d4_tf.paragraphs[0]
    d4_p.text = "👤 LEAD DEVELOPER: MEMBER 4"
    d4_p.font.size = Pt(12)
    d4_p.font.bold = True
    d4_p.font.color.rgb = PRIMARY

    dp4 = d4_tf.add_paragraph()
    dp4.text = "Student Name: [MEMBER 4 NAME]\nStudent ID: [MEMBER 4 ID]\nContribution Weight: 20%"
    dp4.font.size = Pt(12)
    dp4.font.color.rgb = DARK

    dp4 = d4_tf.add_paragraph()
    dp4.text = "\nDedicated Files:\n• src/features/SellerReviewsModal.tsx\n• src/pages/ProfilePage.tsx\n\nGit Branch: feature/seller-reviews"
    dp4.font.size = Pt(10.5)
    dp4.font.color.rgb = TEXT_MUTED

    r4 = s9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(5.3), Inches(1.8), Inches(7.2), Inches(4.8))
    r4.fill.solid()
    r4.fill.fore_color.rgb = CARD_BG
    r4.line.color.rgb = BORDER
    r4_tf = r4.text_frame
    r4_tf.word_wrap = True

    rp4 = r4_tf.paragraphs[0]
    rp4.text = "Key Functional & Technical Deliverables"
    rp4.font.size = Pt(15)
    rp4.font.bold = True
    rp4.font.color.rgb = DARK

    rp_body4 = [
        ("1. Email/Password Authentication:", "Built full sign-up and sign-in with mode-switching, input validation, and Firestore profile synchronization."),
        ("2. Seller Reviews & Trust Badges:", "Created SellerReviewsModal.tsx featuring 5-star ratings, verified badges ('Verified Student', 'Fast Responder'), and satisfaction bars."),
        ("3. Critical Bug Resolution:", "Discovered and resolved the native Android touch lock bug in Expo Go by replacing stacked native Modals with single-modal view routing."),
        ("4. Seamless 'Back to Listing' Nav:", "Provided 60fps instant transitions between listing details and review submissions.")
    ]
    for h, b in rp_body4:
        p_item = r4_tf.add_paragraph()
        p_item.text = f"\n• {h} {b}"
        p_item.font.size = Pt(11)
        p_item.font.color.rgb = TEXT_DARK

    # ==========================================================
    # SLIDE 10: Member 5 Module Deep Dive
    # ==========================================================
    s10 = prs.slides.add_slide(blank_layout)
    set_slide_background(s10)
    add_header(s10, "Module 5: Data Models & Content Moderation Shield")

    dev5 = s10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(4.2), Inches(4.8))
    dev5.fill.solid()
    dev5.fill.fore_color.rgb = TINT_GREEN
    dev5.line.color.rgb = PRIMARY
    d5_tf = dev5.text_frame
    d5_tf.word_wrap = True
    d5_p = d5_tf.paragraphs[0]
    d5_p.text = "👤 LEAD DEVELOPER: MEMBER 5"
    d5_p.font.size = Pt(12)
    d5_p.font.bold = True
    d5_p.font.color.rgb = PRIMARY

    dp5 = d5_tf.add_paragraph()
    dp5.text = "Student Name: [MEMBER 5 NAME]\nStudent ID: [MEMBER 5 ID]\nContribution Weight: 20%"
    dp5.font.size = Pt(12)
    dp5.font.color.rgb = DARK

    dp5 = d5_tf.add_paragraph()
    dp5.text = "\nDedicated Files:\n• src/features/ReportListingModal.tsx\n• src/pages/SavedPage.tsx\n• src/types.ts\n\nGit Branch: feature/report-moderation"
    dp5.font.size = Pt(10.5)
    dp5.font.color.rgb = TEXT_MUTED

    r5 = s10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(5.3), Inches(1.8), Inches(7.2), Inches(4.8))
    r5.fill.solid()
    r5.fill.fore_color.rgb = CARD_BG
    r5.line.color.rgb = BORDER
    r5_tf = r5.text_frame
    r5_tf.word_wrap = True

    rp5 = r5_tf.paragraphs[0]
    rp5.text = "Key Functional & Technical Deliverables"
    rp5.font.size = Pt(15)
    rp5.font.bold = True
    rp5.font.color.rgb = DARK

    rp_body5 = [
        ("1. TypeScript Data Architecture:", "Architected types.ts declaring strict shared models across all 5 developers, eliminating runtime data mismatches."),
        ("2. Saved Shortlist Synchronization:", "Maintained real-time Firestore saved-items subcollection sync with instant optimistic UI toggling."),
        ("3. Content Moderation Shield:", "Built ReportListingModal.tsx allowing confidential reporting of fraudulent items, stolen campus gear, and exam leaks."),
        ("4. Structured Categories:", "Pre-categorized report triggers for admin review (Scam attempt, bait price, copyright infringement, duplicate spam).")
    ]
    for h, b in rp_body5:
        p_item = r5_tf.add_paragraph()
        p_item.text = f"\n• {h} {b}"
        p_item.font.size = Pt(11)
        p_item.font.color.rgb = TEXT_DARK

    # ==========================================================
    # SLIDE 11: Evolution from Base Template (Before vs After)
    # ==========================================================
    s11 = prs.slides.add_slide(blank_layout)
    set_slide_background(s11)
    add_header(s11, "Project Evolution: Initial Prototype vs Final Engineering")

    comp_table_data = [
        ("FEATURE AREA", "INITIAL PROTOTYPE", "OUR FINAL MARKETPLACE", "KEY ENGINEERING IMPACT"),
        ("Campus & Search", "Generic 'North Campus' strings", "17 Sri Lankan State Universities & Price Filters", "Hyper-localized campus trading"),
        ("Messaging", "Empty 20-line stub (no chat)", "Live Firestore Chat + Make an Offer System", "Real-time bargaining & negotiation"),
        ("Seller Inventory", "No edit/delete, dummy images", "Listing CRUD + Firebase Cloud Storage Uploads", "Complete seller inventory lifecycle"),
        ("Authentication", "Anonymous guest logins only", "Email/Password with Firestore User Sync", "Accountability & student verification"),
        ("Trust & Safety", "None (high scam vulnerability)", "5-Star Reviews, Safe Meetups & Scam Shield", "End-to-end student safety"),
    ]

    t11 = s11.shapes.add_table(len(comp_table_data), 4, Inches(0.8), Inches(1.8), Inches(11.733), Inches(4.5))
    tbl11 = t11.table
    tbl11.columns[0].width = Inches(2.2)
    tbl11.columns[1].width = Inches(2.8)
    tbl11.columns[2].width = Inches(3.6)
    tbl11.columns[3].width = Inches(3.133)

    for r_idx, row in enumerate(comp_table_data):
        for c_idx, val in enumerate(row):
            cell = tbl11.cell(r_idx, c_idx)
            cell.text = val
            cell_p = cell.text_frame.paragraphs[0]
            if r_idx == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = PRIMARY
                cell_p.font.bold = True
                cell_p.font.size = Pt(11)
                cell_p.font.color.rgb = WHITE
            else:
                cell.fill.solid()
                cell.fill.fore_color.rgb = CARD_BG if r_idx % 2 == 1 else TINT_GREEN
                cell_p.font.size = Pt(10)
                cell_p.font.color.rgb = TEXT_DARK
                if c_idx == 2:
                    cell_p.font.bold = True
                    cell_p.font.color.rgb = PRIMARY

    # ==========================================================
    # SLIDE 12: Conclusion & Q&A
    # ==========================================================
    s12 = prs.slides.add_slide(blank_layout)
    set_slide_background(s12, DARK)

    c12 = s12.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.5), Inches(1.2), Inches(10.333), Inches(5.1))
    c12.fill.solid()
    c12.fill.fore_color.rgb = PRIMARY
    c12.line.color.rgb = RGBColor(70, 150, 125)
    c12.line.width = Pt(1.5)

    c12_tf = c12.text_frame
    c12_tf.word_wrap = True

    c12_p1 = c12_tf.paragraphs[0]
    c12_p1.text = "🎓 CAMPUS MARKETPLACE"
    c12_p1.font.size = Pt(14)
    c12_p1.font.bold = True
    c12_p1.font.color.rgb = GOLD
    c12_p1.alignment = PP_ALIGN.CENTER

    c12_p2 = c12_tf.add_paragraph()
    c12_p2.text = "Thank You!"
    c12_p2.font.size = Pt(44)
    c12_p2.font.bold = True
    c12_p2.font.color.rgb = WHITE
    c12_p2.alignment = PP_ALIGN.CENTER

    c12_p3 = c12_tf.add_paragraph()
    c12_p3.text = "\nOur team of 5 delivered an enterprise-grade mobile trading application with 5 dedicated subsystems.\nWe are now delighted to answer any questions from the evaluation panel."
    c12_p3.font.size = Pt(14)
    c12_p3.font.color.rgb = RGBColor(220, 240, 230)
    c12_p3.alignment = PP_ALIGN.CENTER

    c12_p4 = c12_tf.add_paragraph()
    c12_p4.text = "\n[Member 1 Name] · [Member 2 Name] · [Member 3 Name] · [Member 4 Name] · [Member 5 Name]"
    c12_p4.font.size = Pt(12)
    c12_p4.font.bold = True
    c12_p4.font.color.rgb = RGBColor(165, 220, 195)
    c12_p4.alignment = PP_ALIGN.CENTER

    prs.save(output_path)
    print(f"Presentation saved to: {output_path}")

if __name__ == "__main__":
    out_dir = r"d:\my projects\campus-marketplace"
    pptx_path = os.path.join(out_dir, "Campus_Marketplace_Presentation.pptx")
    build_presentation(pptx_path)
